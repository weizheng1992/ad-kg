"""达人库端点（spec §3 + §13.1 数据闭环）。

可见性：Creator 是平台级主数据（schema 注记）——类白名单即可见，不走
row_filter 锚点收敛（ae/optimizer 的锚是案例子图，达人库是选人入口）；
client / creative 类白名单不含 Creator 即 403。合作历史（collabs_with）
仍按 plan.visible_entity_ids 收敛——锚点外的案例不出现在达人详情里。
带货商品（promotes → Product）另过类白名单：无 Product 权限的角色不返回该段。

数据闭环（§13.1）：录入（conversation/20）与 CSV 导入（official_doc/60）两条
入口，与平台抓取（platform_api/80）、种子（system_of_record/100）构成四来源
梯度；入库走同一权威仲裁——更高权威在场不覆盖、冲突去重进 review_queue。
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone

import psycopg
from psycopg.rows import dict_row
from fastapi import APIRouter, Header, HTTPException, Request, UploadFile

from adkg.guards.http import get_graph, resolve_identity
from adkg.retrieval.metrics import STAT_FAMILY, agg_expr

router = APIRouter()

CREATE_CREATOR_ROLES = {"ae", "management", "admin"}
SOURCE_LABEL = {"system_of_record": "种子", "platform_api": "平台抓取",
                "official_doc": "表格导入", "conversation": "手工录入"}


def _collab_stats(env: dict, plan, case_ids: list[str]) -> dict | None:
    """达人合作案例维度的汇总投放统计（spec §3）：unified_metrics 汇总比（§六
    铁律）+ 指标列白名单——与 /graph stats、/brands 同一口径与同一道闸。"""
    import psycopg
    from psycopg.rows import dict_row

    if not case_ids:
        return None
    predicate = f" AND {plan.sql_predicate}" if plan.sql_predicate else ""
    try:
        proj = ", ".join(f"{agg_expr(c)} AS {c}" for c in STAT_FAMILY)
        with psycopg.connect(env["ADKG_PG_RO_DSN"], options="-c statement_timeout=3000",
                             row_factory=dict_row) as conn, conn.cursor() as cur:
            cur.execute(
                f"SELECT {proj}, string_agg(DISTINCT platform::text, ',') AS platforms,"
                " count(DISTINCT date) AS days, min(date)::text AS first_day,"
                " max(date)::text AS last_day FROM unified_metrics"
                " WHERE NOT unmatched AND case_id = ANY(%(ids)s)" + predicate,
                {"ids": case_ids})
            row = cur.fetchone()
    except psycopg.Error:
        return None                                          # 统计失败不阻塞详情（§7.2 降级）
    if not row or row["days"] == 0:
        return None
    return {k: v for k, v in row.items()
            if k not in STAT_FAMILY or k in plan.allowed_columns}


def _require_creator(plan) -> None:
    if plan.deny_all:
        raise HTTPException(403, "行权限未生效（缺有效身份），拒绝读取")
    allowed = set(plan.metadata_filter.get("entity_class__in", []))
    if "Creator" not in allowed:
        raise HTTPException(403, "类 Creator 不在当前角色白名单")


def _next_creator_id(graph) -> str:
    """新达人号：现有 crNN 序号最大值 +1（Creator:cr61…；种子命名法，§12.5）。"""
    with graph.session() as s:
        nums = [int(m.group(1)) for m in
                (re.match(r"Creator:cr(\d+)$", r["id"]) for r in s.run(
                    "MATCH (c:Entity {class: 'Creator'}) RETURN c.id AS id").data())
                if m]
    return f"Creator:cr{max(nums, default=0) + 1:02d}"


def _parse_creator_row(row: dict) -> tuple[str, dict]:
    """表单/CSV 行 → (creator_id|'', props)；必填与数值域校验，失败抛 ValueError。"""
    name = str(row.get("name") or "").strip()
    platform = str(row.get("platform") or "").strip()
    niches = str(row.get("niches") or "").strip()
    if not name or not niches:
        raise ValueError("name 与 niches（视频方向）必填")
    if platform not in ("douyin", "xhs"):
        raise ValueError("platform 只能是 douyin / xhs")
    props: dict = {"name": name, "platform": platform, "niches": niches}
    for key in ("followers", "quote_20s", "quote_image", "gmv_30d", "rating"):
        raw = str(row.get(key) or "").strip()
        if not raw:
            continue
        try:
            v = float(raw)
        except ValueError:
            raise ValueError(f"{key} 不是数字：{raw!r}") from None
        props[key] = round(v, 1) if key == "rating" else int(v)
    city = str(row.get("city") or "").strip()
    if city:
        props["city"] = city
    creator_id = str(row.get("creator_id") or "").strip()
    if creator_id and not re.fullmatch(r"Creator:[A-Za-z0-9_-]+", creator_id):
        raise ValueError(f"creator_id 形态应为 Creator:<slug>（§12.5）：{creator_id!r}")
    return creator_id, props


def record_quote_history(env: dict, creator_id: str, quote_20s, source_type: str,
                         source_id: str) -> None:
    """报价快照（spec §15）：三条入库口共用——历史即事实，只追加。"""
    import psycopg

    if not quote_20s:
        return
    with psycopg.connect(env["ADKG_PG_DSN"]) as c, c.cursor() as cur:
        cur.execute(
            "INSERT INTO creator_quote_history (creator_id, date, quote_20s, source_type,"
            " source_id) VALUES (%s, current_date, %s, %s, %s)"
            " ON CONFLICT (creator_id, date, source_type) DO NOTHING",
            (creator_id, int(quote_20s), source_type, source_id))


def _ingest_creator(env: dict, graph, creator_id: str, props: dict,
                    source_type: str, source_id: str,
                    role: str, person: str | None) -> dict:
    """入库 + 权威仲裁（§13.1，与平台抓取同一规则）：更高权威在场不覆盖、
    冲突去重进 review_queue；返回 {outcome, creator_id}。"""
    import json as _json

    import psycopg
    from adkg.guards.audit import audit
    from adkg.models import Entity, Provenance

    prov = Provenance(source_type=source_type, source_id=source_id,      # type: ignore[arg-type]
                      fetched_at=datetime.now(timezone.utc))
    with graph.session() as s:
        rec = s.run("MATCH (n:Entity {id: $id}) RETURN n.authority AS auth",
                    id=creator_id).single()
    existing = rec["auth"] if rec else None
    if existing is not None and existing > (prov.authority or 0):
        with psycopg.connect(env["ADKG_PG_DSN"]) as c, c.cursor() as cur:
            cur.execute(
                "INSERT INTO review_queue (kind, payload)"
                " SELECT 'conflict', %(payload)s"
                " WHERE NOT EXISTS (SELECT 1 FROM review_queue"
                "  WHERE kind = 'conflict' AND status = 'open'"
                "  AND payload->>'entity' = %(entity)s"
                "  AND payload->>'incoming_source' = %(source)s)",
                {"payload": _json.dumps(
                    {"entity": creator_id, "incoming_source": prov.source_id,
                     "incoming_source_type": source_type,
                     "incoming_authority": prov.authority,
                     "existing_authority": existing,
                     "incoming_props": props}, ensure_ascii=False),
                 "entity": creator_id, "source": prov.source_id})
        audit(env, role, person, "create_creator", creator_id, "deny",
              f"更高权威（{existing}）在场未覆盖，冲突进 review_queue（§12.2）")
        return {"outcome": "conflict", "creator_id": creator_id}
    graph.upsert_entities([Entity(entity_class="Creator", entity_id=creator_id,
                                  props=props, provenance=prov)])
    record_quote_history(env, creator_id, props.get("quote_20s"),
                         source_type, prov.source_id)          # 报价快照（§15）
    audit(env, role, person, "create_creator", creator_id, "allow",
          f"来源 {SOURCE_LABEL.get(source_type, source_type)}（{prov.source_id}）")
    return {"outcome": "created" if existing is None else "merged", "creator_id": creator_id}


@router.post("/creators")
def create_creator(request: Request, body: dict,
                   role: str | None = None, user: str | None = None,
                   authorization: str | None = Header(default=None)) -> dict:
    """达人录入（spec §13.1）：conversation/20 低权威——平台抓到更准数据可覆盖。"""
    role, user, planner = resolve_identity(request, authorization, role=role, user=user)
    if role not in CREATE_CREATOR_ROLES:
        raise HTTPException(403, "达人录入仅 ae / management / admin 可用（spec §13.1）")
    try:
        creator_id, props = _parse_creator_row(body)
    except ValueError as e:
        raise HTTPException(422, str(e)) from None
    if not creator_id:
        creator_id = _next_creator_id(get_graph(request))
    return _ingest_creator(request.app.state.env, get_graph(request), creator_id, props,
                           "conversation", f"form:{creator_id}:{datetime.now():%Y%m%d}",
                           role, user)


@router.post("/creators/import")
async def import_creators(request: Request, file: UploadFile,
                          role: str | None = None, user: str | None = None,
                          authorization: str | None = Header(default=None)) -> dict:
    """CSV / xlsx 批量导入（spec §13.1 + §15）：official_doc/60；xlsx 经 openpyxl
    转行后走与 CSV 完全相同的校验/仲裁/审计路；行级容错不中断整批。"""
    role, user, planner = resolve_identity(request, authorization, role=role, user=user)
    if role not in CREATE_CREATOR_ROLES:
        raise HTTPException(403, "达人导入仅 ae / management / admin 可用（spec §13.1）")
    raw = await file.read()
    if (file.filename or "").lower().endswith((".xlsx", ".xls")):
        try:
            import io as _io

            from openpyxl import load_workbook
            wb = load_workbook(_io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            rows_it = ws.iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(rows_it, [])]
            records = [dict(zip(header, ("" if v is None else v for v in row)))
                       for row in rows_it]
        except Exception as e:                              # noqa: BLE001 解析失败整批拒绝
            raise HTTPException(422, f"xlsx 解析失败：{type(e).__name__}: {e}") from None
        if not header or "name" not in header:
            raise HTTPException(422, "xlsx 缺 name 列（模板列见达人库导入提示）")
    else:
        records = csv.DictReader(io.StringIO(raw.decode("utf-8-sig", errors="replace")))
    batch = f"import:{file.filename}:{datetime.now():%Y%m%d%H%M%S}"
    graph = get_graph(request)
    env = request.app.state.env
    created = merged = conflicts = 0
    errors: list[dict] = []
    line_start = 2
    for line_no, row in enumerate(records, start=line_start):
        try:
            creator_id, props = _parse_creator_row(row)
            if not creator_id:
                creator_id = _next_creator_id(graph)
            out = _ingest_creator(env, graph, creator_id, props, "official_doc", batch,
                                  role, user)
        except ValueError as e:
            errors.append({"line": line_no, "error": str(e)})
            continue
        created += out["outcome"] == "created"
        merged += out["outcome"] == "merged"
        conflicts += out["outcome"] == "conflict"
    return {"created": created, "merged": merged, "conflicts": conflicts,
            "errors": errors[:20], "total": created + merged + conflicts + len(errors)}


# 排序白名单（防注入）——coalesce 兜 NULL 在前（Neo4j DESC 空值会排到头部）
_SORTS = {"followers": "coalesce(n.followers, -1)", "gmv": "coalesce(n.gmv_30d, -1)",
          "quote": "coalesce(n.quote_20s, -1)"}


@router.get("/creators")
def list_creators(request: Request, q: str = "", platform: str = "", niche: str = "",
                  source: str = "", min_followers: int = 0, max_quote: int = 0,
                  sort: str = "followers", page: int = 1, size: int = 20,
                  role: str | None = None, user: str | None = None,
                  authorization: str | None = Header(default=None)) -> dict:
    role, user, planner = resolve_identity(request, authorization, role=role, user=user)
    plan = planner.plan(role, person=user)
    _require_creator(plan)
    size = max(1, min(size, 100))
    page = max(1, page)
    order = _SORTS.get(sort, _SORTS["followers"])
    where = ("($q = '' OR n.name CONTAINS $q)"
             " AND ($platform = '' OR n.platform = $platform)"
             " AND ($niche = '' OR n.niches CONTAINS $niche)"
             " AND ($source = '' OR n.prov_source = $source)"
             " AND coalesce(n.followers, 0) >= $min_followers"
             " AND ($max_quote <= 0 OR coalesce(n.quote_20s, 0) <= $max_quote)")
    args = {"q": q.strip(), "platform": platform.strip(), "niche": niche.strip(),
            "source": source.strip(), "min_followers": min_followers,
            "max_quote": max_quote, "skip": (page - 1) * size, "size": size}
    with get_graph(request).session() as s:
        items = s.run(
            f"MATCH (n:Entity {{class: 'Creator'}}) WHERE {where} "
            "OPTIONAL MATCH (c:Entity)-[:collabs_with]->(n) "
            "WITH n, count(DISTINCT c) AS cases "
            f"ORDER BY {order} DESC SKIP $skip LIMIT $size "
            "RETURN n.id AS id, n.name AS name, n.platform AS platform, n.niches AS niches, "
            "n.followers AS followers, n.quote_20s AS quote_20s, n.quote_image AS quote_image, "
            "n.gmv_30d AS gmv_30d, n.city AS city, n.rating AS rating, cases, "
            "n.prov_source AS prov_source, toString(n.prov_fetched_at) AS fetched_at",
            **args).data()
        total = s.run(f"MATCH (n:Entity {{class: 'Creator'}}) WHERE {where} "
                      "RETURN count(n) AS total", **args).single()["total"]
    return {"items": [{"id": it["id"], "name": it["name"], "platform": it["platform"],
                       "niches": (it["niches"] or "").split("|") if it["niches"] else [],
                       "followers": it["followers"], "quote_20s": it["quote_20s"],
                       "quote_image": it["quote_image"], "gmv_30d": it["gmv_30d"],
                       "city": it["city"], "rating": it["rating"], "cases": it["cases"],
                       "source": SOURCE_LABEL.get(it["prov_source"] or "", it["prov_source"]),
                       "fetched_at": it["fetched_at"]} for it in items],
            "total": total, "page": page, "size": size}


@router.get("/creators/{cid}")
def creator_detail(request: Request, cid: str,
                   role: str | None = None, user: str | None = None,
                   authorization: str | None = Header(default=None)) -> dict:
    role, user, planner = resolve_identity(request, authorization, role=role, user=user)
    plan = planner.plan(role, person=user)
    _require_creator(plan)
    visible = plan.visible_entity_ids
    with get_graph(request).session() as s:
        rec = s.run("MATCH (n:Entity {id: $id}) RETURN n.id AS id, n.name AS name, "
                    "n.platform AS platform, n.niches AS niches, n.followers AS followers, "
                    "n.quote_20s AS quote_20s, n.quote_image AS quote_image, "
                    "n.gmv_30d AS gmv_30d, n.city AS city, n.rating AS rating, "
                    "n.class AS class, n.prov_source AS prov_source, "
                    "n.prov_source_id AS prov_source_id, "
                    "toString(n.prov_fetched_at) AS fetched_at, "
                    "n.authority AS authority", id=cid).single()
        if not rec or rec["class"] != "Creator":
            raise HTTPException(404, f"达人 {cid} 不存在")
        collabs = s.run(
            "MATCH (c:Entity)-[:collabs_with]->(n:Entity {id: $id}) "
            "WHERE c.class = 'Case' "
            "RETURN c.id AS id, c.name AS name, c.stage AS stage, "
            "c.lost_reason AS lost_reason ORDER BY c.id", id=cid).data()
        products = []
        if "Product" in set(plan.metadata_filter.get("entity_class__in", [])):
            products = s.run(
                "MATCH (n:Entity {id: $id})-[:promotes]->(p:Entity) "
                "RETURN p.id AS id, p.name AS name, p.price AS price, "
                "p.spec AS spec, p.target_age AS target_age ORDER BY p.id", id=cid).data()
    visible_cases = [c for c in collabs if not visible or c["id"] in visible]
    if plan.allow_properties:                               # 外部角色属性白名单（§12.4）
        visible_cases = [{"id": c["id"], "name": c["name"], "stage": c["stage"]}
                         for c in visible_cases]
    stats = _collab_stats(request.app.state.env, plan, [c["id"] for c in visible_cases])
    # 报价历史（spec §15）：种子轨迹 + 三条入库口的运行时快照
    history: list[dict] = []
    try:
        with psycopg.connect(request.app.state.env["ADKG_PG_DSN"],
                             row_factory=dict_row) as conn, conn.cursor() as cur:
            cur.execute("SELECT date::text AS date, quote_20s, source_type"
                        " FROM creator_quote_history WHERE creator_id = %(id)s"
                        " ORDER BY date", {"id": cid})
            history = [dict(r) for r in cur.fetchall()]
    except psycopg.Error:
        history = []
    # 合作行指标（spec §14：M10 承诺漏做补上）——与 /cases 列表同源批量统计
    from adkg.api.routes.cases import _case_stats_batch
    case_metrics = _case_stats_batch(request.app.state.env, plan,
                                     [c["id"] for c in visible_cases])
    for c in visible_cases:
        m = case_metrics.get(c["id"]) or {}
        c["cost"] = m.get("cost")
        c["cpa"] = m.get("cpa")
    return {
        "creator": {"id": rec["id"], "name": rec["name"], "platform": rec["platform"],
                    "niches": (rec["niches"] or "").split("|") if rec["niches"] else [],
                    "followers": rec["followers"], "quote_20s": rec["quote_20s"],
                    "quote_image": rec["quote_image"], "gmv_30d": rec["gmv_30d"],
                    "city": rec["city"], "rating": rec["rating"]},
        # 行过滤：锚点角色只见自己子图内的合作（spec §3）；无行过滤角色（management/finance）全见
            "collabs": visible_cases,
            "collabs_filtered": len(collabs) - len(visible_cases),
            "products": products,
            "stats": stats,
            # 来源卡（§13.1 四梯度：provenance 即来源）
            "provenance": {"source": SOURCE_LABEL.get(rec["prov_source"] or "",
                                                      rec["prov_source"]),
                           "source_type": rec["prov_source"],
                           "source_id": rec["prov_source_id"],
                           "fetched_at": rec["fetched_at"],
                           "authority": rec["authority"]},
            "quote_history": history,
    }
